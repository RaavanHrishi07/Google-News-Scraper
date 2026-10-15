
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
import xml.etree.ElementTree as ET

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill, Alignment


NEWS_RSS_URL = "https://news.google.com/rss/search"
REQUEST_TIMEOUT = 15


def get_google_news_result(term, count):
    """Fetch news titles, links, sources, and publication dates."""
    params = {
        "q": term,
        "hl": "en-IN",
        "gl": "IN",
        "ceid": "IN:en",
    }

    try:
        response = requests.get(
            NEWS_RSS_URL,
            params=params,
            timeout=REQUEST_TIMEOUT,
            headers={"User-Agent": "GoogleNewsScraper/1.0"},
        )
        response.raise_for_status()
        root = ET.fromstring(response.content)

    except requests.RequestException as exc:
        print(f"Network error while fetching news: {exc}")
        return []

    except ET.ParseError:
        print("Could not read the news feed. Please try again.")
        return []

    articles = []
    seen_links = set()

    for item in root.findall("./channel/item"):
        title = item.findtext("title", default="").strip()
        link = item.findtext("link", default="").strip()
        published = item.findtext("pubDate", default="").strip()

        source_element = item.find("source")
        source = (
            source_element.text.strip()
            if source_element is not None and source_element.text
            else ""
        )

        if not source and link:
            source = urlparse(link).netloc.removeprefix("www.")

        if title and link and link not in seen_links:
            articles.append(
                {
                    "title": title,
                    "source": source or "Unknown",
                    "published_date": published or "Not available",
                    "link": link,
                }
            )
            seen_links.add(link)

        if len(articles) >= count:
            break

    return articles


def create_output_filename(term, extension):
    """Create a filesystem-safe output filename."""
    safe_term = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", term)
    safe_term = safe_term.strip(" ._") or "news"
    return f"{safe_term[:80]}_news_scrapper.{extension}"


def style_header(worksheet):
    """Apply consistent formatting to a worksheet header."""
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 25


def save_articles_to_excel(articles, output_file, keyword, generated_at):
    """Save a summary sheet and a formatted news articles sheet."""
    article_df = pd.DataFrame(
        articles,
        columns=["title", "source", "published_date", "link"],
    )

    article_df.columns = [
        "Title",
        "News Source",
        "Publication Date",
        "Article Link",
    ]

    summary_df = pd.DataFrame(
        [
            {"Field": "Search Keyword", "Value": keyword},
            {"Field": "Total Articles", "Value": len(articles)},
            {"Field": "Report Generated At", "Value": generated_at},
        ]
    )

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(
            writer,
            index=False,
            sheet_name="Summary",
        )
        article_df.to_excel(
            writer,
            index=False,
            sheet_name="News Articles",
        )

        summary_ws = writer.sheets["Summary"]
        summary_ws.freeze_panes = "A2"
        summary_ws.auto_filter.ref = summary_ws.dimensions
        summary_ws.column_dimensions["A"].width = 25
        summary_ws.column_dimensions["B"].width = 40
        style_header(summary_ws)

        articles_ws = writer.sheets["News Articles"]
        articles_ws.freeze_panes = "A2"
        articles_ws.auto_filter.ref = articles_ws.dimensions
        style_header(articles_ws)

        column_widths = {
            "A": 60,
            "B": 28,
            "C": 35,
            "D": 55,
        }

        for column, width in column_widths.items():
            articles_ws.column_dimensions[column].width = width

        for row in range(2, articles_ws.max_row + 1):
            articles_ws.cell(row=row, column=1).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            for column in (2, 3):
                articles_ws.cell(row=row, column=column).alignment = (
                    Alignment(vertical="top", wrap_text=True)
                )

            link_cell = articles_ws.cell(row=row, column=4)
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"
            link_cell.alignment = Alignment(vertical="top")

            articles_ws.row_dimensions[row].height = 30


def main():
    term = input("Enter the news title keyword: ").strip()

    if not term:
        print("Please enter a search keyword.")
        return

    try:
        count = int(input("Enter the number of articles: ").strip())

        if count < 1:
            print("Article count must be at least 1.")
            return

    except ValueError:
        print("Please enter a valid whole number.")
        return

    print(f"\nSearching Google News for: {term}")

    articles = get_google_news_result(term, count)

    if not articles:
        print("No articles found, or the news feed could not be retrieved.")
        return

    generated_at = datetime.now().astimezone().strftime(
        "%Y-%m-%d %H:%M:%S %Z"
    )

    excel_file = create_output_filename(term, "xlsx")
    csv_file = create_output_filename(term, "csv")

    try:
        save_articles_to_excel(
            articles,
            excel_file,
            term,
            generated_at,
        )

        csv_df = pd.DataFrame(
            articles,
            columns=["title", "source", "published_date", "link"],
        )
        csv_df.columns = [
            "Title",
            "News Source",
            "Publication Date",
            "Article Link",
        ]
        csv_df.to_csv(csv_file, index=False, encoding="utf-8-sig")

    except (OSError, PermissionError, ValueError) as exc:
        print(f"Could not save the report: {exc}")
        return

    print(f"\nSuccessfully saved {len(articles)} articles.")
    print(f"Report generated at: {generated_at}")
    print(f"Excel file: {Path(excel_file).resolve()}")
    print(f"CSV file: {Path(csv_file).resolve()}")


if __name__ == "__main__":
    main()
